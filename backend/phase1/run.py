import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import csv
import tempfile
import pyodbc
import pandas as pd
import openpyxl

from config_loader import load_config


def run(log_fn):
    config = load_config()

    conn_str = (
        f"Driver={{SQL Server}};"
        f"Server={config['sql_server']};"
        f"Database={config['sql_database']};"
        f"Trusted_Connection=yes;"
    )

    log_fn("Reading SQL script...")
    with open(config["sql_query_path"], "r") as f:
        query = f.read()

    log_fn("Connecting to SQL Server...")
    conn = pyodbc.connect(conn_str)
    df = pd.read_sql_query(query, conn)
    conn.close()
    log_fn(f"{len(df)} rows returned from query")

    df = df.apply(lambda col: pd.to_numeric(col, errors="ignore"))

    tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False)
    tmp_path = tmp.name
    tmp.close()
    df.to_csv(tmp_path, index=False)

    log_fn("Writing to Excel import sheet...")
    wb = openpyxl.load_workbook(config["excel_workbook_path"])

    if "import" in wb.sheetnames:
        del wb["import"]

    ws = wb.create_sheet("import")
    with open(tmp_path, "r") as f:
        for row in csv.reader(f):
            ws.append(row)

    wb.save(config["excel_workbook_path"])
    os.unlink(tmp_path)

    log_fn(f"Phase 1 complete — {len(df)} rows imported")
