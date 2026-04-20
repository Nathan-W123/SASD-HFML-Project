import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

import json
import openpyxl

from config_loader import load_config

# TODO: verify these column names against the actual workbook
ML_ID_COL   = "WONUM"       # unique ML identifier column header
FREQ_COL    = "FREQUENCY"   # frequency value column header

TEMP_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "temp", "added_mls.json")


def run(log_fn):
    config = load_config()

    wb = openpyxl.load_workbook(config["excel_workbook_path"])
    ws_main   = wb[config["excel_working_sheet"]]
    ws_import = wb["import"]

    # Read headers from both sheets
    main_headers   = [cell.value for cell in next(ws_main.iter_rows(min_row=1, max_row=1))]
    import_headers = [cell.value for cell in next(ws_import.iter_rows(min_row=1, max_row=1))]

    # TODO: confirm ML_ID_COL and FREQ_COL exist in headers before proceeding
    main_id_idx   = main_headers.index(ML_ID_COL)
    import_id_idx = import_headers.index(ML_ID_COL)
    main_freq_idx   = main_headers.index(FREQ_COL)
    import_freq_idx = import_headers.index(FREQ_COL)

    # Build lookup of existing MLs: id -> row index
    existing = {}
    for i, row in enumerate(ws_main.iter_rows(min_row=2, values_only=True), start=2):
        existing[row[main_id_idx]] = i

    # Build lookup of import MLs: id -> full row
    incoming = {}
    for row in ws_import.iter_rows(min_row=2, values_only=True):
        incoming[row[import_id_idx]] = row

    # Detect added MLs (in import, not in mainlines)
    added_ids = [ml_id for ml_id in incoming if ml_id not in existing]
    log_fn(f"{len(added_ids)} new MLs detected")

    # Append added MLs to mainlines sheet
    for ml_id in added_ids:
        row_data = incoming[ml_id]
        # TODO: confirm column mapping between import sheet and mainlines sheet
        ws_main.append(list(row_data))

    # Update frequency for existing MLs where it has changed
    freq_changes = 0
    for ml_id, row_idx in existing.items():
        if ml_id in incoming:
            new_freq = incoming[ml_id][import_freq_idx]
            current_freq = ws_main.cell(row=row_idx, column=main_freq_idx + 1).value
            if new_freq != current_freq:
                ws_main.cell(row=row_idx, column=main_freq_idx + 1).value = new_freq
                freq_changes += 1

    log_fn(f"{freq_changes} frequency changes updated")

    # Remove import sheet
    del wb["import"]

    wb.save(config["excel_workbook_path"])

    # Write added MLs to temp file for Phase 4 and Phase 5
    os.makedirs(os.path.dirname(TEMP_PATH), exist_ok=True)
    with open(TEMP_PATH, "w") as f:
        json.dump(added_ids, f)

    log_fn(f"Phase 2 complete — {len(added_ids)} MLs added, {freq_changes} frequencies updated")
